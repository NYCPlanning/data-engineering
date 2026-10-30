"""Export the core and product tables to one Excel workbook for sharing.

Run after `dbt build`, which writes the tables and target/manifest.json:

    python scripts/export_excel.py
"""

import json
import re
from datetime import date
from pathlib import Path

import duckdb
import pandas as pd
from openpyxl.chart import BarChart, Reference, ScatterChart, Series
from openpyxl.chart.marker import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

PROJECT_DIR = Path(__file__).parent.parent
DATABASE = PROJECT_DIR / "data" / "analysis_growth.duckdb"
MANIFEST = PROJECT_DIR / "target" / "manifest.json"
# dated by export day, so earlier exports aren't overwritten
OUTPUT = (
    PROJECT_DIR
    / "data"
    / "exports"
    / f"Community District Growth {date.today():%Y%m%d}.xlsx"
)

# Excel caps sheet names at 31 characters
CHARTS_SHEET = "Charts"
DISTRICTS_SHEET = "Community districts"
GROWTH_SHEET = "Population vs housing growth"
COMPONENTS_SHEET = "Where population change went"
WIDE_SHEET = "Housing units by year"
LONG_SHEET = "Housing units by year (long)"
SHEET_MODELS = {
    DISTRICTS_SHEET: "dim_community_district",
    GROWTH_SHEET: "population_vs_housing_growth",
    COMPONENTS_SHEET: "population_change_components",
    LONG_SHEET: "mart_community_district_date",
}


# the notebook's growth gap colors
RED = "B2182B"
BLUE = "2166AC"

# product models describe their columns as "Same as <model>.<column>."
SAME_AS = re.compile(r"^Same as (\w+)\.(\w+)\.$")


def model_node(manifest: dict, model: str) -> dict:
    return manifest["nodes"][f"model.analysis_growth.{model}"]


def source_column(manifest: dict, column: dict) -> dict:
    """The column's docs, following a "Same as" pointer to where they're written."""
    if same_as := SAME_AS.match(column["description"].strip()):
        model, name = same_as.groups()
        return model_node(manifest, model)["columns"][name]
    return column


def display_names(manifest: dict, model: str) -> dict[str, str]:
    """Column name -> display name, from each column's meta.display_name in the dbt YAML."""
    names = {}
    for column in model_node(manifest, model)["columns"].values():
        meta = source_column(manifest, column)["config"].get("meta", {})
        if "display_name" in meta:
            names[column["name"]] = meta["display_name"]
    return names


def model_docs(manifest: dict, sheet: str, model: str, renamed: bool) -> list[dict]:
    node = model_node(manifest, model)
    names = display_names(manifest, model) if renamed else {}
    rows = [
        {
            "sheet": sheet,
            "column": "",
            "dbt column": "",
            "description": node["description"],
        }
    ]
    for column in node["columns"].values():
        if column["name"] != "geom":
            rows.append(
                {
                    "sheet": sheet,
                    "column": names.get(column["name"], column["name"]),
                    "dbt column": column["name"],
                    "description": source_column(manifest, column)["description"],
                }
            )
    return rows


# openpyxl can't autofit, so widths come from the longest value in each column. Longer text,
# like the column descriptions, wraps at MAX_WIDTH instead of widening the column.
MAX_WIDTH = 80


def autofit_columns(worksheet, df: pd.DataFrame) -> None:
    for index, name in enumerate(df.columns, start=1):
        longest = max([len(str(name)), *df[name].astype(str).str.len()])
        letter = get_column_letter(index)
        worksheet.column_dimensions[letter].width = min(longest + 2, MAX_WIDTH)
        # dates have no time part, so the default "yyyy-mm-dd h:mm:ss" would show only zeros
        # and overflow the fitted width
        if pd.api.types.is_datetime64_any_dtype(df[name]):
            for (cell,) in worksheet.iter_rows(min_col=index, max_col=index, min_row=2):
                cell.number_format = "yyyy-mm-dd"
        if longest > MAX_WIDTH:
            for (cell,) in worksheet.iter_rows(min_col=index, max_col=index, min_row=2):
                cell.alignment = Alignment(wrap_text=True, vertical="top")


def style_chart(chart) -> None:
    # one color per series, not one per point
    chart.varyColors = False
    # titles beside the plot, not over it
    chart.title.overlay = False
    for axis in [chart.x_axis, chart.y_axis]:
        # openpyxl hides axes unless told otherwise
        axis.delete = False
        # tick labels at the chart's edge, not where the axes cross 0 inside the plot
        axis.tickLblPos = "low"
        if axis.title is not None:
            axis.title.overlay = False
        if axis.majorGridlines is not None:
            axis.majorGridlines.spPr = GraphicalProperties(
                ln=LineProperties(solidFill="D9D9D9")
            )
    if isinstance(chart, BarChart) and chart.type == "bar":
        # ranked bars flip the district axis, which turns "low" into the top edge. The value
        # axis already sits at the bottom, so its labels go next to it.
        chart.y_axis.tickLblPos = "nextTo"


def ranked_bars(title: str, height: float) -> BarChart:
    """Horizontal bars, one per district, in sheet order from the top."""
    bars = BarChart()
    bars.type = "bar"
    bars.title = title
    # categories run bottom to top by default; reversed, the first row is at the top. A value
    # axis crossing at the last category then sits at the bottom. (crosses on an axis sets where
    # the other axis meets it, so it goes on the value axis.)
    bars.x_axis.scaling.orientation = "maxMin"
    bars.y_axis.crosses = "max"
    bars.gapWidth = 40
    bars.width, bars.height = 18, height
    return bars


def sheet_column(worksheet, df: pd.DataFrame, header: str, with_header: bool = False):
    index = df.columns.get_loc(header) + 1
    return Reference(
        worksheet, min_col=index, min_row=1 if with_header else 2, max_row=len(df) + 1
    )


def growth_charts(worksheet, growth: pd.DataFrame, names: dict[str, str]) -> list:
    """Charts of the growth sheet's data, for the charts sheet."""

    def column(name: str, header: bool = False) -> Reference:
        return sheet_column(worksheet, growth, names[name], with_header=header)

    scatter = ScatterChart()
    scatter.title = "Population vs housing growth, 2010 to 2020"
    scatter.x_axis.title = "Units built (%)"
    scatter.y_axis.title = "Population change (%)"
    points = Series(
        column("population_change_2010_2020_pct"),
        column("units_built_2010_2020_pct"),
    )
    points.marker.symbol = "circle"
    points.marker.size = 6
    points.marker.graphicalProperties = GraphicalProperties(solidFill=BLUE)
    points.marker.graphicalProperties.line.solidFill = BLUE
    # markers only: a scatter chart connects its points by default
    points.graphicalProperties.line.noFill = True
    scatter.series.append(points)
    scatter.legend = None
    scatter.width, scatter.height = 18, 12

    # bars in sheet order, which is sorted by the gap, largest first
    gap_column = "population_vs_housing_growth_gap_2010_2020_pts"
    bars = ranked_bars(
        "Population vs housing growth gap, 2010 to 2020 (pts)", height=30
    )
    bars.add_data(column(gap_column, header=True), titles_from_data=True)
    bars.set_categories(column("cdta_name"))
    # red where population grew faster than housing, blue where housing grew faster
    for index, gap in enumerate(growth[names[gap_column]]):
        bars.series[0].dPt.append(
            # Excel ignores a point's color without the two flags it always writes itself
            DataPoint(
                idx=index,
                invertIfNegative=False,
                bubble3D=False,
                spPr=GraphicalProperties(solidFill=RED if gap > 0 else BLUE),
            )
        )
    bars.legend = None

    for chart in [scatter, bars]:
        style_chart(chart)
    return [scatter, bars]


# component -> (legend label, color), the same colors as the notebook's chart and map
COMPONENTS = {
    "population_change_from_units_built_2010_2020_pts": ("Units built", BLUE),
    "population_change_from_occupied_units_not_built_2010_2020_pts": (
        "Occupied units not built",
        "E08214",
    ),
    "population_change_from_household_size_2010_2020_pts": ("Household size", RED),
    "population_change_from_group_quarters_2010_2020_pts": ("Group quarters", "6A51A3"),
}


def components_chart(
    worksheet, components: pd.DataFrame, names: dict[str, str]
) -> BarChart:
    """Stacked bars of where each district's population change went, in sheet order."""
    bars = ranked_bars(
        "Where population change went, 2010 to 2020 (pts of 2010 population)", height=30
    )
    bars.grouping = "stacked"
    bars.overlap = 100
    for name, (label, color) in COMPONENTS.items():
        series = Series(sheet_column(worksheet, components, names[name]), title=label)
        series.graphicalProperties = GraphicalProperties(solidFill=color)
        bars.series.append(series)
    bars.set_categories(sheet_column(worksheet, components, names["cdta_name"]))
    # the value axis labels leave no room for a legend across the top
    bars.legend.position = "r"
    style_chart(bars)
    return bars


def column_docs(manifest: dict) -> pd.DataFrame:
    rows = model_docs(manifest, GROWTH_SHEET, SHEET_MODELS[GROWTH_SHEET], renamed=True)
    rows += model_docs(
        manifest, COMPONENTS_SHEET, SHEET_MODELS[COMPONENTS_SHEET], renamed=True
    )
    rows += model_docs(
        manifest, DISTRICTS_SHEET, SHEET_MODELS[DISTRICTS_SHEET], renamed=True
    )
    rows += [
        {
            "sheet": WIDE_SHEET,
            "column": "",
            "description": f"The same values as {LONG_SHEET}, with one row per community "
            "district and one column per as_of_date.",
        },
        {
            "sheet": WIDE_SHEET,
            "column": "borocd",
            "description": "Community district.",
        },
        {
            "sheet": WIDE_SHEET,
            "column": "community_district_name",
            "description": "Community district, e.g. Brooklyn Community District 10.",
        },
        {
            "sheet": WIDE_SHEET,
            "column": "cdta_name",
            "description": "Code and name of the district's CDTA, e.g. BK10 Bay Ridge-Dyker Heights.",
        },
        {
            "sheet": WIDE_SHEET,
            "column": "<dates>",
            "description": f"housing_units_from_2020_census as of that date. See {LONG_SHEET}.",
        },
    ]
    rows += model_docs(manifest, LONG_SHEET, SHEET_MODELS[LONG_SHEET], renamed=False)
    docs = pd.DataFrame(rows)
    docs["dbt column"] = docs["dbt column"].fillna(docs["column"])
    docs["description"] = docs["description"].str.split().str.join(" ")
    return docs.rename(
        columns={"sheet": "Sheet", "column": "Column", "description": "Description"}
    )


def main() -> None:
    manifest = json.loads(MANIFEST.read_text())
    with duckdb.connect(str(DATABASE), read_only=True) as con:
        # geometry doesn't fit in a spreadsheet cell
        districts = (
            con.sql("SELECT * FROM dim_community_district").df().drop(columns="geom")
        )
        growth = con.sql("SELECT * FROM population_vs_housing_growth").df()
        components = con.sql("SELECT * FROM population_change_components").df()
        housing_long = con.sql("SELECT * FROM mart_community_district_date").df()
    district_names = display_names(manifest, SHEET_MODELS[DISTRICTS_SHEET])
    growth_names = display_names(manifest, SHEET_MODELS[GROWTH_SHEET])
    components_names = display_names(manifest, SHEET_MODELS[COMPONENTS_SHEET])
    housing_wide = housing_long.assign(
        as_of_date=housing_long["as_of_date"].dt.strftime("%Y-%m-%d")
    ).pivot(
        index=["borocd", "community_district_name", "cdta_name"],
        columns="as_of_date",
        values="housing_units_from_2020_census",
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(OUTPUT) as writer:
        for sheet, df in [
            (GROWTH_SHEET, growth.rename(columns=growth_names)),
            (COMPONENTS_SHEET, components.rename(columns=components_names)),
            (DISTRICTS_SHEET, districts.rename(columns=district_names)),
            (WIDE_SHEET, housing_wide.reset_index()),
            (LONG_SHEET, housing_long),
            ("Column descriptions", column_docs(manifest)),
        ]:
            df.to_excel(writer, sheet_name=sheet, index=False, freeze_panes=(1, 0))
            autofit_columns(writer.sheets[sheet], df)
            print(f"{sheet}: {len(df)} rows")
        charts = writer.book.create_sheet(CHARTS_SHEET, 0)
        for chart, anchor in zip(
            growth_charts(
                writer.sheets[GROWTH_SHEET],
                growth.rename(columns=growth_names),
                growth_names,
            ),
            ["B2", "M2"],
        ):
            charts.add_chart(chart, anchor)
        charts.add_chart(
            components_chart(
                writer.sheets[COMPONENTS_SHEET],
                components.rename(columns=components_names),
                components_names,
            ),
            "X2",
        )
    print(f"wrote {OUTPUT.relative_to(PROJECT_DIR)}")


if __name__ == "__main__":
    main()
